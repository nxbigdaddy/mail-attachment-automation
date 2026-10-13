from docx import Document
from openpyxl import load_workbook
import os

base_dir = r"<workspace>\mail-auto\attachments_20260531_132532"

print("=" * 60)
print("Reading 示例员工's files content...")
print("=" * 60)

dinghai_files = [
    "示例员工/48_工作日报-示例员工-26.5.19.docx",
    "示例员工/60_出差日志2026.5.20（无锡派克新材）.docx",
    "示例员工/69_工作日报-示例员工-26.5.20&5.21.docx",
    "示例员工/69_工作日报-示例员工-26.5.22.docx",
    "示例员工/72_工作日报-示例员工-26.5.25.docx",
    "示例员工/74_工作日报-示例员工-26.5.26.docx",
]

for filepath in dinghai_files:
    full_path = os.path.join(base_dir, filepath)
    if os.path.exists(full_path):
        print(f"\n📄 {filepath}")
        print("-" * 60)
        try:
            doc = Document(full_path)
            for para in doc.paragraphs[:20]:
                if para.text.strip():
                    print(para.text)
        except Exception as e:
            print(f"Error: {e}")

print("\n" + "=" * 60)
print("Reading 水彬's files content...")
print("=" * 60)

shuibin_files = [
    "水彬/49_2026-05-19 工作日报.xlsx",
    "水彬/57_2026-05-20 工作日报.xlsx",
    "水彬/61_2026-05-21 工作日报.xlsx",
    "水彬/68_2026-05-22 工作日报.xlsx",
    "水彬/73_2026-05-25 工作日报.xlsx",
    "水彬/75_2026-05-26 工作日报.xlsx",
]

for filepath in shuibin_files:
    full_path = os.path.join(base_dir, filepath)
    if os.path.exists(full_path):
        print(f"\n📄 {filepath}")
        print("-" * 60)
        try:
            wb = load_workbook(full_path)
            sheet = wb.active
            for row in sheet.iter_rows(values_only=True, max_row=30):
                row_str = [str(cell) if cell else "" for cell in row]
                if any(cell.strip() for cell in row_str):
                    print("\t".join(row_str))
        except Exception as e:
            print(f"Error: {e}")
