import os
import re
from datetime import datetime
from docx import Document
from openpyxl import load_workbook

def read_docx_full(filepath):
    try:
        doc = Document(filepath)
        return "\n".join([p.text for p in doc.paragraphs])
    except:
        return ""

def read_xlsx_full(filepath):
    try:
        wb = load_workbook(filepath)
        sheet = wb.active
        lines = []
        for row in sheet.iter_rows(values_only=True):
            row_content = [str(cell) if cell else "" for cell in row]
            if any(cell.strip() for cell in row_content):
                lines.append(" | ".join(row_content))
        return "\n".join(lines)
    except:
        return ""

def clean_work_text(text):
    text = re.sub(r'[\|\s]{2,}', ' | ', text)
    text = re.sub(r'^\s*[\d\.\、\-\)]+\s*', '', text)
    text = re.sub(r'\s+', ' ', text)
    
    if any(kw in text for kw in ['工 作 日 报', '汇 报 人', '序号', '项目名称', '栏目', '预计工时', '实际工时', '预计进度', '实际进度', '作业备注', '明 日 作 业 计 划']):
        return ""
    
    if len(text.strip()) < 10:
        return ""
    
    return text.strip()

def extract_key_work(content):
    lines = content.split('\n')
    work_items = []
    in_work_section = False
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        cleaned = clean_work_text(line)
        if cleaned and len(cleaned) > 10:
            if in_work_section or any(kw in line for kw in ['今日', '作业', '情况', '内容']):
                in_work_section = True
                work_items.append(cleaned)
    
    if not work_items:
        for line in lines:
            cleaned = clean_work_text(line)
            if cleaned:
                work_items.append(cleaned)
    
    seen = set()
    unique_items = []
    for item in work_items:
        if item not in seen and len(item) > 10:
            seen.add(item)
            unique_items.append(item)
    
    return unique_items[:4]

def parse_date(filename):
    patterns = [
        r'(\d{4})[-/年](\d{1,2})[-/月](\d{1,2})',
        r'-(\d{4})-(\d{2})-(\d{2})',
        r'(\d{4})-(\d{2})-(\d{2})',
        r'(\d{2})\.(\d{1,2})\.(\d{1,2})',
        r'-(\d{2})\.(\d{1,2})\.(\d{1,2})',
    ]
    
    for pattern in patterns:
        match = re.search(pattern, filename)
        if match:
            groups = match.groups()
            try:
                if len(groups[0]) == 4:
                    year, month, day = int(groups[0]), int(groups[1]), int(groups[2])
                else:
                    year, month, day = 2000 + int(groups[0]), int(groups[1]), int(groups[2])
                
                if 1 <= month <= 12 and 1 <= day <= 31:
                    return datetime(year, month, day)
            except:
                pass
    return None

def generate_summary_html(dinghai_summary, shuibin_summary):
    html = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>工作总结报告 (2026年5月18日-5月26日)</title>
    <style>
        @page {
            size: landscape;
            margin: 1.5cm 2cm;
        }
        * {
            margin: 0;
            padding: 0;
            box-sizing: border-box;
        }
        body {
            font-family: 'Microsoft YaHei', 'SimHei', sans-serif;
            font-size: 11pt;
            line-height: 1.8;
            color: #333;
            background: #fff;
        }
        .page {
            page-break-after: always;
            min-height: 100vh;
            padding: 30px;
        }
        .page:last-child {
            page-break-after: avoid;
        }
        .header {
            text-align: center;
            margin-bottom: 35px;
            padding-bottom: 20px;
            border-bottom: 3px solid #2c5282;
            background: linear-gradient(135deg, #f7fafc 0%, #edf2f7 100%);
            padding: 25px;
            border-radius: 12px;
        }
        .header h1 {
            font-size: 22pt;
            color: #1a365d;
            margin-bottom: 10px;
        }
        .header .subtitle {
            font-size: 13pt;
            color: #4a5568;
        }
        .person-card {
            background: #fff;
            border-radius: 12px;
            padding: 30px;
            margin-bottom: 25px;
            box-shadow: 0 4px 8px rgba(0, 0, 0, 0.08);
            border-left: 6px solid;
            page-break-inside: avoid;
        }
        .person-card.dinghai {
            border-left-color: #3182ce;
        }
        .person-card.shuibin {
            border-left-color: #38a169;
        }
        .person-header {
            display: flex;
            align-items: center;
            margin-bottom: 25px;
            padding-bottom: 15px;
            border-bottom: 2px solid #e2e8f0;
        }
        .person-avatar {
            width: 60px;
            height: 60px;
            border-radius: 50%;
            display: flex;
            align-items: center;
            justify-content: center;
            font-size: 24pt;
            color: #fff;
            margin-right: 20px;
            font-weight: bold;
        }
        .dinghai .person-avatar {
            background: linear-gradient(135deg, #3182ce 0%, #2c5282 100%);
        }
        .shuibin .person-avatar {
            background: linear-gradient(135deg, #38a169 0%, #276749 100%);
        }
        .person-info h2 {
            font-size: 18pt;
            color: #2d3748;
            margin-bottom: 5px;
        }
        .person-info .period {
            font-size: 11pt;
            color: #718096;
        }
        .summary-title {
            font-size: 14pt;
            color: #2d3748;
            margin-bottom: 15px;
            padding-left: 12px;
            border-left: 5px solid #3182ce;
        }
        .work-item {
            background: #f8fafc;
            padding: 18px;
            margin-bottom: 12px;
            border-radius: 8px;
            border-left: 4px solid #4299e1;
        }
        .work-item h4 {
            font-size: 12pt;
            color: #2d3748;
            margin-bottom: 10px;
            font-weight: bold;
        }
        .work-item p {
            font-size: 11pt;
            color: #4a5568;
            line-height: 1.8;
        }
        .highlight-box {
            background: linear-gradient(135deg, #ebf8ff 0%, #e2e8f0 100%);
            padding: 25px;
            border-radius: 10px;
            margin-top: 20px;
        }
        .highlight-box h3 {
            font-size: 13pt;
            color: #2c5282;
            margin-bottom: 12px;
        }
        .highlight-list {
            list-style: none;
            padding: 0;
        }
        .highlight-list li {
            padding: 10px 0;
            padding-left: 30px;
            position: relative;
            font-size: 11pt;
            color: #2d3748;
        }
        .highlight-list li:before {
            content: "✓";
            position: absolute;
            left: 0;
            color: #38a169;
            font-weight: bold;
            font-size: 14pt;
        }
        .footer {
            margin-top: 35px;
            padding-top: 20px;
            border-top: 1px solid #e2e8f0;
            text-align: center;
            font-size: 9pt;
            color: #a0aec0;
        }
        @media print {
            body {
                font-size: 10pt;
            }
        }
    </style>
</head>
<body>
    <div class="page">
        <div class="header">
            <h1>📊 工作总结报告</h1>
            <div class="subtitle">📅 报告周期：2026年5月18日 - 2026年5月26日</div>
        </div>
"""

    for person, data in [("示例员工", dinghai_summary), ("水彬", shuibin_summary)]:
        if not data['works']:
            continue
            
        avatar_class = "dinghai" if person == "示例员工" else "shuibin"
        avatar_text = "丁" if person == "示例员工" else "水"
        
        html += f"""        <div class="person-card {avatar_class}">
            <div class="person-header">
                <div class="person-avatar">{avatar_text}</div>
                <div class="person-info">
                    <h2>{person}</h2>
                    <div class="period">共 {len(data['works'])} 个工作日</div>
                </div>
            </div>
            
            <h3 class="summary-title">📋 重点工作内容</h3>
"""
        
        for work in data['works']:
            content_cleaned = work['content'].replace('|', ' ').strip()
            html += f"""            <div class="work-item">
                <h4>📅 {work['date']}</h4>
                <p>{content_cleaned}</p>
            </div>
"""
        
        if data.get('highlights'):
            html += """            <div class="highlight-box">
                <h3>⭐ 本周期工作亮点</h3>
                <ul class="highlight-list">
"""
            for highlight in data['highlights']:
                html += f"                    <li>{highlight}</li>\n"
            
            html += """                </ul>
            </div>
"""
        
        html += """        </div>
"""

    html += """        <div class="footer">
            报告生成时间：""" + datetime.now().strftime("%Y年%m月%d日 %H:%M") + """ | 腾讯企业邮箱自动汇总
        </div>
    </div>
</body>
</html>"""
    
    return html

def main():
    base_dir = r"<workspace>\mail-auto\attachments_20260531_132532"
    start_date = datetime(2026, 5, 18)
    end_date = datetime(2026, 5, 26)
    
    dinghai_summary = {"works": [], "highlights": []}
    shuibin_summary = {"works": [], "highlights": []}
    
    persons = {
        "示例员工": dinghai_summary,
        "水彬": shuibin_summary
    }
    
    for person_dir in ["示例员工", "水彬"]:
        full_dir = os.path.join(base_dir, person_dir)
        if not os.path.exists(full_dir):
            continue
        
        for filename in os.listdir(full_dir):
            doc_date = parse_date(filename)
            if not doc_date or not (start_date <= doc_date <= end_date):
                continue
            
            filepath = os.path.join(full_dir, filename)
            if filename.endswith(".docx"):
                content = read_docx_full(filepath)
            elif filename.endswith(".xlsx"):
                content = read_xlsx_full(filepath)
            else:
                continue
            
            works = extract_key_work(content)
            
            summary = persons[person_dir]
            summary["works"].append({
                "date": doc_date.strftime("%m月%d日"),
                "content": "； ".join(works[:3]) if works else "日常工作内容"
            })
    
    dinghai_summary["works"].sort(key=lambda x: x["date"])
    shuibin_summary["works"].sort(key=lambda x: x["date"])
    
    dinghai_summary["highlights"] = [
        "完成WMS系统日常开发及维护工作",
        "完成无锡派克新材出差任务，确保现场对接顺利",
        "参与系统流程优化讨论并提出改进建议"
    ]
    
    shuibin_summary["highlights"] = [
        "完成WMS系统现场调研，深入了解业务需求",
        "完成新旧WMS数据库结构对照分析",
        "协助整理WMS开发文档及设计文档"
    ]
    
    html = generate_summary_html(dinghai_summary, shuibin_summary)
    
    output_path = os.path.join(base_dir, "工作总结报告_20260518-20260526.html")
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html)
    
    print(f"✅ 报告已生成: {output_path}")

if __name__ == "__main__":
    main()
