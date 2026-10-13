from docx import Document
import os

base_dir = r"<workspace>\mail-auto\attachments_20260531_132532"

print("搜索'客户拜访'关键词出现的文件：")
print("=" * 60)

for filename in os.listdir(os.path.join(base_dir, "示例员工")):
    if filename.endswith((".docx", ".doc")):
        filepath = os.path.join(base_dir, "示例员工", filename)
        try:
            if filename.endswith(".docx"):
                doc = Document(filepath)
                content = "\n".join([p.text for p in doc.paragraphs])
            else:
                doc = Document(filepath)
                content = "\n".join([p.text for p in doc.paragraphs])

            if "客户拜访" in content:
                print(f"\n📄 文件: {filename}")
                print("-" * 60)
                for line in content.split('\n'):
                    if "客户拜访" in line:
                        print(f"  {line.strip()}")
        except Exception as e:
            print(f"Error reading {filename}: {e}")
