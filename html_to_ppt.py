from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import re
import os
from datetime import datetime

def html_to_ppt(html_path, ppt_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    with open(html_path, 'r', encoding='utf-8') as f:
        html_content = f.read()

    title_match = re.search(r'<title>(.*?)</title>', html_content)
    main_title = title_match.group(1) if title_match else "下属工作汇总报告"

    pages = html_content.split('<div class="page">')

    for i, page_content in enumerate(pages[1:], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        if 'cover' in page_content:
            title_match = re.search(r'<h1[^>]*>(.*?)</h1>', page_content)
            subtitle_match = re.search(r'<div class="subtitle">(.*?)</div>', page_content)
            meta_match = re.search(r'<div class="meta">(.*?)</div>', page_content)

            if title_match:
                title_text = clean_html(title_match.group(1))
                add_text_box(slide, title_text, Inches(0.5), Inches(2.5), Inches(12.333), Inches(1.5),
                           font_size=44, bold=True, color=RGBColor(26, 54, 93), align=PP_ALIGN.CENTER)

            if subtitle_match:
                subtitle_text = clean_html(subtitle_match.group(1))
                add_text_box(slide, subtitle_text, Inches(0.5), Inches(4.5), Inches(12.333), Inches(1),
                           font_size=24, color=RGBColor(74, 85, 104), align=PP_ALIGN.CENTER)

            if meta_match:
                meta_text = clean_html(meta_match.group(1))
                add_text_box(slide, meta_text, Inches(0.5), Inches(5.5), Inches(12.333), Inches(0.8),
                           font_size=16, color=RGBColor(113, 128, 150), align=PP_ALIGN.CENTER)

        elif 'main-content' in page_content:
            add_text_box(slide, "下属工作汇总报告", Inches(0.3), Inches(0.3), Inches(12.733), Inches(0.8),
                        font_size=28, bold=True, color=RGBColor(26, 54, 93), align=PP_ALIGN.CENTER)

            person_sections = page_content.split('<div class="person-section">')
            y_offset = 1.3

            for j, section in enumerate(person_sections[1:], 1):
                if y_offset > 6.5:
                    break

                name_match = re.search(r'<h2>(.*?)</h2>', section)
                days_match = re.search(r'<div class="days">(.*?)</div>', section)

                if name_match:
                    name = clean_html(name_match.group(1))
                    days = clean_html(days_match.group(1)) if days_match else ""

                    add_text_box(slide, f"{name}  {days}", Inches(0.5), Inches(y_offset), Inches(12.333), Inches(0.5),
                               font_size=20, bold=True, color=RGBColor(44, 55, 72))
                    y_offset += 0.5

                    work_items = re.findall(r'<div class="work-item">(.*?)</div>', section)
                    for item in work_items[:4]:
                        if y_offset > 6.3:
                            break
                        work_text = clean_html(item)
                        add_shape_rect(slide, Inches(0.7), Inches(y_offset), Inches(11.933), Inches(0.7),
                                     RGBColor(247, 250, 252))
                        add_text_box(slide, work_text, Inches(0.9), Inches(y_offset + 0.1), Inches(11.533), Inches(0.5),
                                   font_size=12, color=RGBColor(74, 85, 104))
                        y_offset += 0.8

        elif 'highlights-section' in page_content:
            add_text_box(slide, "⭐ 工作重点成果", Inches(0.3), Inches(0.3), Inches(12.733), Inches(0.8),
                        font_size=28, bold=True, color=RGBColor(26, 54, 93), align=PP_ALIGN.CENTER)

            highlight_cards = page_content.split('<div class="highlight-card">')
            x_positions = [0.5, 6.9]

            for j, card in enumerate(highlight_cards[1:3], 0):
                name_match = re.search(r'<h3[^>]*>(.*?)</h3>', card)
                items = re.findall(r'<li>(.*?)</li>', card)

                if name_match:
                    name = clean_html(name_match.group(1))
                    add_shape_rect(slide, Inches(x_positions[j]), Inches(1.3), Inches(5.9), Inches(5.7),
                                 RGBColor(235, 248, 255))

                    add_text_box(slide, name, Inches(x_positions[j] + 0.3), Inches(1.5), Inches(5.3), Inches(0.5),
                               font_size=18, bold=True, color=RGBColor(44, 82, 130))

                    y_pos = 2.2
                    for item in items[:4]:
                        item_text = clean_html(item)
                        item_text = item_text.replace('✓', '').replace('strong', '').strip()
                        add_text_box(slide, f"✓ {item_text}", Inches(x_positions[j] + 0.3), Inches(y_pos),
                                   Inches(5.3), Inches(0.9), font_size=12, color=RGBColor(74, 85, 104))
                        y_pos += 1.0

    prs.save(ppt_path)
    print(f"✅ PPT已生成: {ppt_path}")

def clean_html(text):
    text = re.sub(r'<[^>]+>', '', text)
    text = text.replace('&nbsp;', ' ')
    text = text.replace('&lt;', '<')
    text = text.replace('&gt;', '>')
    text = text.replace('&amp;', '&')
    text = text.strip()
    return text

def add_text_box(slide, text, left, top, width, height, font_size=14, bold=False,
                 color=RGBColor(0, 0, 0), align=PP_ALIGN.LEFT):
    textbox = slide.shapes.add_textbox(left, top, width, height)
    tf = textbox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return textbox

def add_shape_rect(slide, left, top, width, height, fill_color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape

if __name__ == "__main__":
    html_file = r"<workspace>\mail-auto\attachments_20260531_132532\管理者汇总报告_20260518-20260526.html"
    ppt_file = r"<workspace>\mail-auto\attachments_20260531_132532\管理者汇总报告_20260518-20260526.pptx"

    html_to_ppt(html_file, ppt_file)
